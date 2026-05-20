"""
3D Technology Services - Project Manager
Flask + SQLite app for tracking BOM, daily installs, material/labor cost.
"""
import csv
import io
import re
import sqlite3
from datetime import datetime
from pathlib import Path

from flask import (
    Flask, abort, flash, redirect, render_template, request, url_for,
)

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

try:
    import pdfplumber
except ImportError:
    pdfplumber = None


BASE_DIR = Path(__file__).resolve().parent
DB_PATH = BASE_DIR / "database.db"
UPLOAD_DIR = BASE_DIR / "uploads"
UPLOAD_DIR.mkdir(exist_ok=True)

ALLOWED_EXT = {".csv", ".xlsx", ".pdf"}

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 16 * 1024 * 1024
app.secret_key = "3dts-pm-dev-key-change-in-prod"


# ---------- Database ----------

def get_db():
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    conn.execute("PRAGMA foreign_keys = ON")
    return conn


def init_db():
    conn = get_db()
    conn.executescript("""
    CREATE TABLE IF NOT EXISTS projects (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        name TEXT NOT NULL,
        pm_name TEXT NOT NULL,
        bid_labor_hours REAL NOT NULL DEFAULT 0,
        status TEXT NOT NULL DEFAULT 'Active',
        file_name TEXT,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP
    );
    CREATE TABLE IF NOT EXISTS bom_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_id INTEGER NOT NULL,
        material TEXT NOT NULL,
        bom_qty REAL NOT NULL DEFAULT 0,
        unit_price REAL NOT NULL DEFAULT 0,
        hours_per_unit REAL NOT NULL DEFAULT 0,
        needs_review INTEGER NOT NULL DEFAULT 0,
        FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS daily_entries (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        project_id INTEGER NOT NULL,
        entry_date TEXT NOT NULL,
        total_hours REAL NOT NULL DEFAULT 0,
        notes TEXT,
        created_at TEXT NOT NULL DEFAULT CURRENT_TIMESTAMP,
        FOREIGN KEY (project_id) REFERENCES projects(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS daily_entry_items (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        daily_entry_id INTEGER NOT NULL,
        bom_item_id INTEGER NOT NULL,
        qty_installed REAL NOT NULL DEFAULT 0,
        FOREIGN KEY (daily_entry_id) REFERENCES daily_entries(id) ON DELETE CASCADE,
        FOREIGN KEY (bom_item_id) REFERENCES bom_items(id) ON DELETE CASCADE
    );
    CREATE TABLE IF NOT EXISTS labor_manual (
        id INTEGER PRIMARY KEY AUTOINCREMENT,
        material_pattern TEXT NOT NULL UNIQUE,
        hours_per_unit REAL NOT NULL DEFAULT 0,
        category TEXT,
        notes TEXT
    );
    """)
    conn.commit()
    conn.close()


# ---------- Helpers ----------

def _norm(s):
    return re.sub(r"[^a-z0-9]+", "", (s or "").lower())


def _to_float(v, default=0.0):
    if v is None or v == "":
        return default
    try:
        return float(re.sub(r"[^0-9.\-]", "", str(v)))
    except (ValueError, TypeError):
        return default


def _find_col(headers, candidates):
    norm_map = {_norm(h): h for h in headers if h}
    for c in candidates:
        nc = _norm(c)
        if nc in norm_map:
            return norm_map[nc]
    for c in candidates:
        nc = _norm(c)
        for h_norm, h in norm_map.items():
            if h_norm and (nc in h_norm or h_norm in nc):
                return h
    return None


def lookup_labor(conn, material):
    """Look up labor hours for a material by exact or fuzzy name match."""
    norm = _norm(material)
    rows = conn.execute(
        "SELECT material_pattern, hours_per_unit FROM labor_manual"
    ).fetchall()
    for r in rows:
        if _norm(r["material_pattern"]) == norm:
            return r["hours_per_unit"]
    for r in rows:
        pn = _norm(r["material_pattern"])
        if pn and (pn in norm or norm in pn):
            return r["hours_per_unit"]
    return None


# ---------- BOM parsers ----------

def _parse_rows(headers, rows):
    mat_col = _find_col(headers, ["material", "description", "item", "part", "product"])
    qty_col = _find_col(headers, ["qty", "quantity", "count"])
    price_col = _find_col(headers, ["unit price", "price", "unit cost", "cost"])
    hours_col = _find_col(headers, ["hours per unit", "hrs per unit", "labor hours", "hours", "hrs", "labor"])

    items = []
    for row in rows:
        if not isinstance(row, dict):
            continue
        material = row.get(mat_col) if mat_col else None
        if material is None or str(material).strip() == "":
            continue
        items.append({
            "material": str(material).strip(),
            "bom_qty": _to_float(row.get(qty_col) if qty_col else None),
            "unit_price": _to_float(row.get(price_col) if price_col else None),
            "hours_per_unit": _to_float(row.get(hours_col) if hours_col else None) if hours_col else None,
        })
    return items


def parse_csv_file(path):
    with open(path, "r", encoding="utf-8-sig", errors="ignore", newline="") as f:
        reader = csv.DictReader(f)
        headers = reader.fieldnames or []
        rows = list(reader)
    return _parse_rows(headers, rows)


def parse_xlsx_file(path):
    if load_workbook is None:
        raise RuntimeError("openpyxl not installed — pip install openpyxl")
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    grid = list(ws.iter_rows(values_only=True))
    if not grid:
        return []
    # Find first non-empty row as header
    header_idx = 0
    for i, r in enumerate(grid):
        if any(c is not None and str(c).strip() for c in r):
            header_idx = i
            break
    headers = [str(c).strip() if c is not None else "" for c in grid[header_idx]]
    rows = []
    for r in grid[header_idx + 1:]:
        rows.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
    return _parse_rows(headers, rows)


def parse_pdf_file(path):
    if pdfplumber is None:
        raise RuntimeError("pdfplumber not installed — pip install pdfplumber")
    items = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages:
            tables = page.extract_tables() or []
            for table in tables:
                if not table or len(table) < 2:
                    continue
                headers = [str(c).strip() if c is not None else "" for c in table[0]]
                rows = []
                for r in table[1:]:
                    rows.append({headers[i]: r[i] for i in range(min(len(headers), len(r)))})
                items.extend(_parse_rows(headers, rows))
    return items


def parse_bom(path, filename):
    ext = Path(filename).suffix.lower()
    if ext == ".csv":
        return parse_csv_file(path)
    if ext == ".xlsx":
        return parse_xlsx_file(path)
    if ext == ".pdf":
        return parse_pdf_file(path)
    raise ValueError(f"Unsupported file type: {ext}")


# ---------- Routes ----------

@app.route("/")
def dashboard():
    conn = get_db()
    projects = conn.execute(
        "SELECT * FROM projects ORDER BY created_at DESC"
    ).fetchall()
    conn.close()
    return render_template("dashboard.html", projects=projects)


@app.route("/projects", methods=["POST"])
def create_project():
    file = request.files.get("bom_file")
    name = (request.form.get("name") or "").strip()
    pm_name = (request.form.get("pm_name") or "").strip()
    bid_hours = _to_float(request.form.get("bid_labor_hours"))

    if not file or not file.filename or not name or not pm_name:
        flash("Missing project name, PM name, or BOM file.", "error")
        return redirect(url_for("dashboard"))

    ext = Path(file.filename).suffix.lower()
    if ext not in ALLOWED_EXT:
        flash(f"Unsupported file type: {ext}. Use .csv, .xlsx, or .pdf.", "error")
        return redirect(url_for("dashboard"))

    safe_name = re.sub(r"[^A-Za-z0-9._-]", "_", file.filename)
    save_path = UPLOAD_DIR / f"{datetime.now().strftime('%Y%m%d%H%M%S')}_{safe_name}"
    file.save(str(save_path))

    try:
        items = parse_bom(save_path, file.filename)
    except Exception as e:
        flash(f"Failed to parse BOM: {e}", "error")
        return redirect(url_for("dashboard"))

    if not items:
        flash("No items found in BOM. Ensure the file has Material/Qty columns.", "error")
        return redirect(url_for("dashboard"))

    conn = get_db()
    cur = conn.cursor()
    cur.execute(
        "INSERT INTO projects (name, pm_name, bid_labor_hours, file_name) VALUES (?, ?, ?, ?)",
        (name, pm_name, bid_hours, file.filename),
    )
    project_id = cur.lastrowid

    for it in items:
        hours = it["hours_per_unit"]
        needs_review = 0
        if hours is None or hours == 0:
            looked_up = lookup_labor(conn, it["material"])
            if looked_up is not None:
                hours = looked_up
            else:
                hours = 0.0
                needs_review = 1
        cur.execute(
            "INSERT INTO bom_items (project_id, material, bom_qty, unit_price, hours_per_unit, needs_review) "
            "VALUES (?, ?, ?, ?, ?, ?)",
            (project_id, it["material"], it["bom_qty"], it["unit_price"], hours, needs_review),
        )

    conn.commit()
    conn.close()
    flash(f"Project '{name}' created with {len(items)} line items.", "success")
    return redirect(url_for("dashboard"))


@app.route("/project/<int:project_id>/delete", methods=["POST"])
def delete_project(project_id):
    conn = get_db()
    conn.execute("DELETE FROM projects WHERE id = ?", (project_id,))
    conn.commit()
    conn.close()
    flash("Project deleted.", "success")
    return redirect(url_for("dashboard"))


@app.route("/project/<int:project_id>/daily", methods=["GET"])
def daily(project_id):
    conn = get_db()
    project = conn.execute(
        "SELECT * FROM projects WHERE id = ?", (project_id,)
    ).fetchone()
    if not project:
        abort(404)

    items = conn.execute("""
        SELECT b.*, COALESCE(SUM(d.qty_installed), 0) AS installed
        FROM bom_items b
        LEFT JOIN daily_entry_items d ON d.bom_item_id = b.id
        WHERE b.project_id = ?
        GROUP BY b.id
        ORDER BY b.id
    """, (project_id,)).fetchall()

    history = conn.execute("""
        SELECT * FROM daily_entries
        WHERE project_id = ?
        ORDER BY entry_date DESC, created_at DESC
    """, (project_id,)).fetchall()

    conn.close()
    return render_template(
        "daily.html",
        project=project,
        items=items,
        history=history,
        today=datetime.now().strftime("%Y-%m-%d"),
    )


@app.route("/project/<int:project_id>/daily", methods=["POST"])
def save_daily(project_id):
    conn = get_db()
    project = conn.execute(
        "SELECT id FROM projects WHERE id = ?", (project_id,)
    ).fetchone()
    if not project:
        abort(404)

    entry_date = request.form.get("entry_date") or datetime.now().strftime("%Y-%m-%d")
    total_hours = _to_float(request.form.get("total_hours"))
    notes = request.form.get("notes", "").strip()

    installs = {}
    for k, v in request.form.items():
        if not k.startswith("installed_"):
            continue
        try:
            bom_id = int(k.split("_", 1)[1])
        except ValueError:
            continue
        qty = _to_float(v)
        if qty > 0:
            installs[bom_id] = qty

    if not installs and total_hours == 0:
        flash("Nothing to save — enter installed quantities or hours.", "error")
        conn.close()
        return redirect(url_for("daily", project_id=project_id))

    cur = conn.cursor()
    cur.execute(
        "INSERT INTO daily_entries (project_id, entry_date, total_hours, notes) VALUES (?, ?, ?, ?)",
        (project_id, entry_date, total_hours, notes),
    )
    entry_id = cur.lastrowid
    for bom_id, qty in installs.items():
        cur.execute(
            "INSERT INTO daily_entry_items (daily_entry_id, bom_item_id, qty_installed) VALUES (?, ?, ?)",
            (entry_id, bom_id, qty),
        )
    conn.commit()
    conn.close()
    flash(f"Saved daily entry for {entry_date} — {len(installs)} item(s), {total_hours:.2f} hours.", "success")
    return redirect(url_for("daily", project_id=project_id))


@app.route("/project/<int:project_id>/admin")
def admin(project_id):
    conn = get_db()
    project = conn.execute(
        "SELECT * FROM projects WHERE id = ?", (project_id,)
    ).fetchone()
    if not project:
        abort(404)

    items = conn.execute("""
        SELECT b.*, COALESCE(SUM(d.qty_installed), 0) AS installed
        FROM bom_items b
        LEFT JOIN daily_entry_items d ON d.bom_item_id = b.id
        WHERE b.project_id = ?
        GROUP BY b.id
        ORDER BY b.id
    """, (project_id,)).fetchall()

    actual_field_hours = conn.execute(
        "SELECT COALESCE(SUM(total_hours), 0) AS h FROM daily_entries WHERE project_id = ?",
        (project_id,),
    ).fetchone()["h"]

    rows = []
    total_material_cost = 0.0
    total_labor_book_hours = 0.0
    total_installed_qty = 0.0
    total_installed_labor_hours = 0.0
    for it in items:
        line_material_cost = it["unit_price"] * it["bom_qty"]
        labor_book_hours = it["hours_per_unit"] * it["bom_qty"]
        installed = it["installed"]
        installed_labor_hours = it["hours_per_unit"] * installed
        remaining_qty = it["bom_qty"] - installed
        remaining_labor_hours = labor_book_hours - installed_labor_hours
        rows.append({
            "id": it["id"],
            "material": it["material"],
            "unit_price": it["unit_price"],
            "bom_qty": it["bom_qty"],
            "hours_per_unit": it["hours_per_unit"],
            "needs_review": it["needs_review"],
            "line_material_cost": line_material_cost,
            "labor_book_hours": labor_book_hours,
            "installed": installed,
            "installed_labor_hours": installed_labor_hours,
            "remaining_qty": remaining_qty,
            "remaining_labor_hours": remaining_labor_hours,
        })
        total_material_cost += line_material_cost
        total_labor_book_hours += labor_book_hours
        total_installed_qty += installed
        total_installed_labor_hours += installed_labor_hours

    bid_hours = project["bid_labor_hours"]
    totals = {
        "material_cost": total_material_cost,
        "labor_book_hours": total_labor_book_hours,
        "installed_qty": total_installed_qty,
        "installed_labor_hours": total_installed_labor_hours,
        "actual_field_hours": actual_field_hours,
        "labor_book_vs_bid": total_labor_book_hours - bid_hours,
        "actual_vs_bid": actual_field_hours - bid_hours,
        "actual_vs_labor_book": actual_field_hours - total_installed_labor_hours,
    }

    conn.close()
    return render_template("admin.html", project=project, rows=rows, totals=totals)


@app.route("/project/<int:project_id>/update", methods=["POST"])
def update_project(project_id):
    bid_hours = _to_float(request.form.get("bid_labor_hours"))
    status = request.form.get("status", "Active")
    conn = get_db()
    conn.execute(
        "UPDATE projects SET bid_labor_hours = ?, status = ? WHERE id = ?",
        (bid_hours, status, project_id),
    )
    conn.commit()
    conn.close()
    flash("Project updated.", "success")
    return redirect(url_for("admin", project_id=project_id))


@app.route("/project/<int:project_id>/bom-item/<int:bom_id>", methods=["POST"])
def update_bom_item(project_id, bom_id):
    hours = _to_float(request.form.get("hours_per_unit"))
    unit_price_raw = request.form.get("unit_price")
    conn = get_db()
    if unit_price_raw not in (None, ""):
        conn.execute(
            "UPDATE bom_items SET hours_per_unit = ?, unit_price = ?, needs_review = 0 "
            "WHERE id = ? AND project_id = ?",
            (hours, _to_float(unit_price_raw), bom_id, project_id),
        )
    else:
        conn.execute(
            "UPDATE bom_items SET hours_per_unit = ?, needs_review = 0 "
            "WHERE id = ? AND project_id = ?",
            (hours, bom_id, project_id),
        )
    conn.commit()
    conn.close()
    flash("Item updated.", "success")
    return redirect(url_for("admin", project_id=project_id))


@app.route("/labor-manual")
def labor_manual():
    conn = get_db()
    rows = conn.execute(
        "SELECT * FROM labor_manual ORDER BY category, material_pattern"
    ).fetchall()
    conn.close()
    return render_template("labor_manual.html", rows=rows)


@app.route("/labor-manual", methods=["POST"])
def save_labor_manual():
    pattern = (request.form.get("material_pattern") or "").strip()
    hours = _to_float(request.form.get("hours_per_unit"))
    category = (request.form.get("category") or "").strip()
    notes = (request.form.get("notes") or "").strip()
    if not pattern:
        flash("Material pattern is required.", "error")
        return redirect(url_for("labor_manual"))
    conn = get_db()
    conn.execute("""
        INSERT INTO labor_manual (material_pattern, hours_per_unit, category, notes)
        VALUES (?, ?, ?, ?)
        ON CONFLICT(material_pattern) DO UPDATE SET
            hours_per_unit = excluded.hours_per_unit,
            category = excluded.category,
            notes = excluded.notes
    """, (pattern, hours, category, notes))
    conn.commit()
    conn.close()
    flash(f"Saved '{pattern}'.", "success")
    return redirect(url_for("labor_manual"))


@app.route("/labor-manual/<int:lm_id>/delete", methods=["POST"])
def delete_labor_manual(lm_id):
    conn = get_db()
    conn.execute("DELETE FROM labor_manual WHERE id = ?", (lm_id,))
    conn.commit()
    conn.close()
    return redirect(url_for("labor_manual"))


init_db()


if __name__ == "__main__":
    app.run(debug=True, port=5000, host="127.0.0.1")
